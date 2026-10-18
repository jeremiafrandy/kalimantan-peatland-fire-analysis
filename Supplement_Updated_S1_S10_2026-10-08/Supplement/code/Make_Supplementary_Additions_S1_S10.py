"""Assemble the author-supplied workflow and Table S10 into an editable supplement.

The workflow PNG is a static author-supplied artwork. It is preserved without
redrawing or interpreting the pathways as evidence of source attribution.
"""
from pathlib import Path
import argparse
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from Make_Table_S10_Rainfall_Category_Counts import CAPTION, NOTE

FIGURE_CAPTION = ('Figure S1. Overview of the observational analysis workflow. '
    'Rainfall and MODIS records were used to examine antecedent rainfall '
    'associations during 2020–2025. Burned-area observations were compared '
    'within and outside mapped peatland in Central Kalimantan during '
    'August–November 2023. Published outdoor PM₂.₅ measurements and atmospheric '
    'reanalysis provided supporting air-quality context.')

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--output',type=Path,default=Path(__file__).resolve().parents[1])
    a=ap.parse_args(); out=a.output
    image=out/'figures/Figure_S1_Observational_Analysis_Workflow.png'
    table=pd.read_csv(out/'tables/Table_S10_Rainfall_Category_Counts.csv')
    captions=out/'Figure_and_Table_Captions.md'
    captions.write_text('# Supplementary figure and table captions\n\n'+
        FIGURE_CAPTION+'\n\n'+
        'Figure S2. Observation-subset sensitivity of the conditional rainfall association. '
        'Points and intervals compare detection-day odds after 0–10 mm with >100 mm of '
        'preceding seven-day rainfall. Models condition on grid cell × year-month and '
        'adjust for earlier rainfall. Intervals are approximate. These checks do not '
        'correct for cloud-related observation availability or validate predictions.\n\n'+
        'Figure S3. Daily outdoor PM₂.₅ during the 2023 sensor campaign. Lines show the '
        'mean across reporting sensor sites, the fixed two-site mean, and the CAMS '
        'daily mean for the shared model cell in UTC. Published sensor date labels are '
        'retained. This comparison provides air-quality context without identifying '
        'the source of measured particles.\n\n'+CAPTION+'\n\n'+NOTE+'\n',encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Caption']:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12)
        s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman')
        s.element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
        for attr in list(s.element.rPr.rFonts.attrib):
            if 'theme' in attr.lower(): del s.element.rPr.rFonts.attrib[attr]
        for attr in ['eastAsia','cs']:s.element.rPr.rFonts.set(qn('w:'+attr),'Times New Roman')
        for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(16)
    doc.styles['Title'].font.bold=True
    doc.styles['Caption'].font.bold=False
    normal=doc.styles['Normal'].paragraph_format
    normal.space_after=Pt(8);normal.line_spacing=1.15
    doc.add_paragraph('Supplementary workflow and rainfall observations',style='Title')
    doc.add_paragraph('The workflow summarizes the complementary observational analyses. '
        'Table S10 reports how grid-days were distributed across the rainfall categories '
        'used in the conditional analysis.')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image),width=Inches(6.87))
    doc.inline_shapes[0]._inline.docPr.set('descr','Three parallel branches for rainfall '
        'associations, mapped burning contrasts, and supporting air-quality context.')
    doc.add_paragraph(FIGURE_CAPTION,style='Caption')
    doc.add_page_break()
    doc.add_paragraph(CAPTION,style='Caption')
    t=doc.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.22,1.83,1.73,2.09]
    headers=['Preceding seven-day rainfall (mm)',
        'Grid-days in the complete analytic panel',
        'Grid-days in informative strata',
        'Recorded detection grid-days in informative strata']
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for _,row in table.iterrows():
        cells=t.add_row().cells
        cells[0].text=str(row.iloc[0])
        for i in range(1,4):cells[i].text=f'{int(row.iloc[i]):,}'
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4')
        x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        cannot_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cannot_split)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);t.columns[j].width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge,value in [('top','100'),('bottom','100'),('left','95'),('right','95')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
            pr.append(marg)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(11)
                    r.bold=(i==0 or i==len(t.rows)-1)
    p=doc.add_paragraph(NOTE);p.paragraph_format.space_before=Pt(10)
    doc.add_paragraph('The binary outcome indicates at least one recorded MODIS detection '
        'classified as presumed vegetation fire (type = 0) on a grid-day. Seven-day '
        'rainfall covers D−7 to D−1. A grid-day with no recorded detection is not '
        'independently verified as a clear-sky, fire-free observation.')
    doc.core_properties.title='Supplementary workflow and rainfall observations'
    doc.core_properties.subject='Figure S1 and Table S10 for the Kalimantan peatland study'
    doc.core_properties.author=''
    path=out/'Supplementary_Additions_S1_S10.docx';doc.save(path)
    # Publication PDF of the supplied artwork on a white page; no raster redraw.
    with Image.open(image) as im: width,height=im.size
    pdf=out/'figures/Figure_S1_Observational_Analysis_Workflow.pdf'
    c=canvas.Canvas(str(pdf),pagesize=(width,height),pageCompression=1)
    c.setTitle('Figure S1 Observational analysis workflow')
    c.setFillColorRGB(1,1,1);c.rect(0,0,width,height,fill=1,stroke=0)
    c.drawImage(ImageReader(str(image)),0,0,width=width,height=height,mask='auto')
    c.showPage();c.save()
    print(path)

if __name__=='__main__':main()
